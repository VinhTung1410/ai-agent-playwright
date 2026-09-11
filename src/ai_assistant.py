import os
import io
import re
import datetime
import logging
from typing import Dict, List, Any, Optional
import dotenv
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

dotenv.load_dotenv()
logger = logging.getLogger(__name__)

# Preferred models in order of priority
MODEL_CANDIDATES = [
    "gemini-3.6-flash",
    "gemini-flash-latest",
    "gemini-3.8-flash",
    "gemini-2.5-flash-lite",
    "gemini-1.5-flash"
]


def get_api_key(custom_key: Optional[str] = None) -> Optional[str]:
    """
    Récupère la clé d'API Gemini (priorité à la clé personnalisée passée en paramètre,
    puis aux fichiers .env à la racine ou dans .venv, puis aux variables d'environnement / st.secrets).
    """
    if custom_key and custom_key.strip():
        return custom_key.strip()

    # Recharger dynamiquement depuis les emplacements possibles (.env à la racine ou dans .venv)
    current_dir = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.abspath(os.path.join(current_dir, ".."))
    
    env_paths = [
        os.path.join(root_dir, ".env"),
        os.path.join(root_dir, ".venv", ".env"),
        os.path.join(current_dir, ".env"),
        ".env",
        ".venv/.env"
    ]
    for p in env_paths:
        if os.path.exists(p) and os.path.isfile(p):
            dotenv.load_dotenv(dotenv_path=p, override=True)
            break

    key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    
    # Fallback pour Streamlit Cloud via st.secrets
    if not key:
        try:
            import streamlit as st
            if hasattr(st, "secrets"):
                for secret_name in ["GEMINI_API_KEY", "GOOGLE_API_KEY", "gemini_api_key", "google_api_key"]:
                    if secret_name in st.secrets:
                        val = st.secrets[secret_name]
                        if val and str(val).strip():
                            key = str(val).strip()
                            break
                if not key and "gemini" in st.secrets:
                    gem_sec = st.secrets["gemini"]
                    if isinstance(gem_sec, dict):
                        key = gem_sec.get("api_key") or gem_sec.get("GEMINI_API_KEY")
        except Exception:
            pass

    return key.strip() if key and key.strip() else None


def call_gemini_api(prompt: str, system_instruction: Optional[str] = None, api_key: Optional[str] = None) -> str:
    """
    Exécute un appel à l'API Gemini avec gestion des exceptions et bascule automatique de modèle.
    """
    key = get_api_key(api_key)
    if not key:
        raise ValueError("Clé API Gemini introuvable. Veuillez configurer GEMINI_API_KEY dans votre fichier .env.")

    # 1. Tentative avec le SDK moderne google-genai
    try:
        from google import genai
        from google.genai import types
        client = genai.Client(api_key=key)
        
        last_err = None
        for model_name in MODEL_CANDIDATES:
            try:
                gen_config = types.GenerateContentConfig(
                    system_instruction=system_instruction,
                    temperature=0.7
                ) if system_instruction else types.GenerateContentConfig(temperature=0.7)
                
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                    config=gen_config
                )
                if response and response.text:
                    return response.text.strip()
            except Exception as e:
                last_err = e
                logger.warning(f"Modèle {model_name} indisponible: {e}. Essai du modèle suivant...")
                continue
        
        if last_err:
            raise last_err

    except ImportError:
        logger.info("google-genai non disponible, tentative avec google.generativeai...")

    # 2. Fallback avec google.generativeai si disponible
    try:
        import google.generativeai as legacy_genai
        legacy_genai.configure(api_key=key)
        for model_name in ["gemini-1.5-flash", "gemini-pro"]:
            try:
                model = legacy_genai.GenerativeModel(
                    model_name=model_name,
                    system_instruction=system_instruction
                )
                response = model.generate_content(prompt)
                if response and response.text:
                    return response.text.strip()
            except Exception:
                continue
    except Exception as e:
        logger.error(f"Échec de l'appel Gemini: {e}")

    raise RuntimeError("Impossible de contacter l'API Gemini avec les modèles configurés.")


def detect_cover_letter_placeholders(text: str) -> List[str]:
    """
    Détecte les balises de substitution non remplies (ex: [Votre Nom], [Adresse], [Date])
    pour avertir le candidat avant soumission.
    """
    if not text:
        return []
    matches = re.findall(r'(\[[A-ZÀ-Ÿa-zà-ÿ0-9_\s\.\-\/\'\"]{3,40}\])', text)
    filtered = []
    for m in matches:
        m_lower = m.lower()
        if any(kw in m_lower for kw in ["votre", "nom", "adresse", "téléphone", "email", "date", "entreprise", "école", "formation"]):
            if m not in filtered:
                filtered.append(m)
    return filtered


def generate_cover_letter(
    cv_text: str,
    job_details: Dict[str, Any],
    matched_skills: List[str],
    missing_skills: List[str],
    candidate_name: str = "Lucas Martin",
    candidate_info: Optional[Dict[str, str]] = None,
    tone: str = "Professionnel & Dynamique",
    api_key: Optional[str] = None
) -> str:
    """
    Génère une lettre de motivation en français, sur-mesure, structurée selon les standards
    de recrutement français (Vous - Moi - Nous) en exploitant l'intégralité du texte du CV.
    """
    company = job_details.get("Company", "l'Entreprise")
    title = job_details.get("Job Title", job_details.get("Title", "Business Analyst"))
    location = job_details.get("Location", "France")
    contract = job_details.get("Employment Type", "Alternance")
    job_desc = job_details.get("Description", "")[:4000]
    
    matched_str = ", ".join(matched_skills) if matched_skills else "Business Analysis, Gestion de projet"
    missing_str = ", ".join(missing_skills) if missing_skills else "Aucune majeure"

    info = candidate_info or {}
    c_name = info.get("name") or candidate_name
    c_email = info.get("email", "")
    c_phone = info.get("phone", "")
    c_loc = info.get("location", "France")
    c_title = info.get("title", "")

    contact_block = f"- Nom complet : {c_name}\n"
    if c_email:
        contact_block += f"- Email : {c_email}\n"
    if c_phone:
        contact_block += f"- Téléphone : {c_phone}\n"
    if c_loc:
        contact_block += f"- Ville : {c_loc}\n"
    if c_title:
        contact_block += f"- Profil actuel : {c_title}\n"

    system_instruction = (
        "Tu es un expert senior en recrutement et coach carrière de premier plan sur le marché du travail en France. "
        "Tu rédiges des lettres de motivation percutantes, personnalisées et parfaitement rédigées en français soutenu "
        "pour des postes de Business Analyst, Product Owner et Data Analyst (Alternance, Stage, CDI). "
        "Tu appliques rigoureusement la structure 'Vous - Moi - Nous', sans phrases creuses ni clichés. "
        "IMPORTANT : Utilise DIRECTEMENT les coordonnées réelles fournies dans l'en-tête, sans jamais laisser de placeholders génériques comme '[Votre Nom]' ou '[Téléphone]'."
    )

    prompt = f"""
Rédige une lettre de motivation complète et ultra-personnalisée pour la candidature suivante :

### INFORMATIONS CANDIDAT :
{contact_block}
- Compétences validées avec l'offre (Points forts) : {matched_str}
- Compétences manquantes ou à valoriser : {missing_str}
- Texte intégral du CV :
\"\"\"{cv_text[:8000]}\"\"\"

### INFORMATIONS OFFRE D'EMPLOI :
- Poste visé : {title}
- Entreprise : {company}
- Localisation : {location}
- Type de contrat : {contract}
- Description du poste :
\"\"\"{job_desc}\"\"\"

### DIRECTIVES DE RÉDACTION :
1. **Ton souhaité :** {tone}.
2. **Structure obligatoire :**
   - **En-tête :** Coordonnées candidat complètes (avec {c_name}, {c_phone}, {c_email}), Coordonnées entreprise ({company}), Date du jour, Objet clair : **Objet : Candidature au poste de {title} ({contract})**.
   - **Introduction (Accroche & 'Vous') :** Montrer une compréhension précise de l'activité de {company} et des enjeux du poste de {title}.
   - **Corps de texte ('Moi') :** Valoriser 2 ou 3 réalisations concrètes du CV qui démontrent la maîtrise de {matched_str}. Mentionner habilement l'aptitude à monter en compétence sur {missing_str}.
   - **Projection ('Nous') :** Expliquer comment le candidat va créer de la valeur concrète dès ses premières semaines dans l'équipe.
   - **Conclusion :** Demande proactive d'entretien et formule de politesse formelle et élégante (ex: "Je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées").
3. **Format :** Rends le texte en Markdown soigné avec sections distinctes, sans placeholders entre crochets.
"""

    return call_gemini_api(prompt, system_instruction=system_instruction, api_key=api_key)


def export_cover_letter_docx(
    cover_letter_markdown: str = "",
    candidate_name: Any = None,
    job_details: Optional[Dict[str, Any]] = None,
    candidate_info: Optional[Dict[str, Any]] = None,
    **kwargs
) -> bytes:
    """
    Génère un document Microsoft Word (.docx) professionnel respectant les standards
    de mise en page des lettres de motivation en France (marges 2.5cm, typographie soignée, en-têtes alignés).
    """
    if not cover_letter_markdown and "cover_letter_text" in kwargs:
        cover_letter_markdown = kwargs["cover_letter_text"]
        
    info = dict(candidate_info or {})
    if isinstance(candidate_name, dict) and not candidate_info:
        info = dict(candidate_name)
    elif isinstance(candidate_name, str) and candidate_name.strip():
        info["name"] = candidate_name.strip()
        
    if not job_details:
        job_details = kwargs.get("job_raw", {})
    candidate_info = info
    doc = docx.Document()
    
    # 1. Marges standard (2.5 cm partout)
    for section in doc.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)

    # Style par défaut (Calibri 11pt, couleur #1E293B)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # 2. Table d'en-tête (2 colonnes sans bordures)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(3.5)

    # Cellule Gauche : Candidat
    cell_cand = table.cell(0, 0)
    p_cand = cell_cand.paragraphs[0]
    p_cand.paragraph_format.line_spacing = 1.15
    p_cand.paragraph_format.space_after = Pt(2)
    
    cand_name = candidate_info.get("name") or "Candidat"
    r_name = p_cand.add_run(cand_name + "\n")
    r_name.bold = True
    r_name.font.size = Pt(12)
    r_name.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    if candidate_info.get("title"):
        r_title = p_cand.add_run(candidate_info["title"] + "\n")
        r_title.italic = True
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        
    cand_details = []
    if candidate_info.get("phone"):
        cand_details.append(f"Tél : {candidate_info['phone']}")
    if candidate_info.get("email"):
        cand_details.append(f"Email : {candidate_info['email']}")
    if candidate_info.get("location"):
        cand_details.append(candidate_info['location'])
        
    if cand_details:
        p_cand.add_run("\n".join(cand_details))

    # Cellule Droite : Entreprise & Date
    cell_comp = table.cell(0, 1)
    p_comp = cell_comp.paragraphs[0]
    p_comp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_comp.paragraph_format.line_spacing = 1.15
    p_comp.paragraph_format.space_after = Pt(2)
    
    comp_name = job_details.get("Company", "l'Entreprise")
    r_comp = p_comp.add_run(comp_name + "\n")
    r_comp.bold = True
    r_comp.font.size = Pt(11)
    
    date_str = datetime.date.today().strftime("%d/%m/%Y")
    city = candidate_info.get("location", "Paris").split(",")[0].strip()
    p_comp.add_run(f"À l'attention du Service Recrutement\n{city}, le {date_str}")

    # Espace après tableau
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(12)

    # 3. Ligne Objet
    title = job_details.get("Job Title", job_details.get("Title", "Business Analyst"))
    contract = job_details.get("Employment Type", "")
    objet_text = f"Objet : Candidature au poste de {title}"
    if contract and contract != "Non spécifié":
        objet_text += f" ({contract})"
        
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(14)
    r_obj = p_obj.add_run(objet_text)
    r_obj.bold = True
    r_obj.font.size = Pt(11)
    r_obj.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # 4. Traitement du corps de texte
    clean_lines = []
    skip_header = True
    for line in cover_letter_markdown.splitlines():
        line_s = line.strip()
        if line_s.lower().startswith("objet") or ("candidature" in line_s.lower() and len(line_s) < 80):
            skip_header = False
            continue
        if skip_header and (line_s.startswith("#") or "madame" in line_s.lower() or "monsieur" in line_s.lower()):
            skip_header = False
        if not skip_header:
            clean_lines.append(line)
            
    body_text = "\n".join(clean_lines).strip()
    if not body_text:
        body_text = cover_letter_markdown

    # Découper par paragraphes
    paragraphs = [p.strip() for p in body_text.split("\n\n") if p.strip()]
    for p_content in paragraphs:
        text_clean = p_content
        text_clean = re.sub(r'^#+\s*', '', text_clean)
        
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        
        # Gestion du gras Markdown (**gras**)
        parts = re.split(r'(\*\*.*?\*\*)', text_clean)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)

    # 5. Bloc signature
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(16)
    r_sig = p_sig.add_run(cand_name)
    r_sig.bold = True

    # 6. Sauvegarder dans un flux binaire
    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream.getvalue()


def generate_inmail_message(
    cv_text: str,
    job_details: Dict[str, Any],
    matched_skills: List[str],
    candidate_name: str = "Lucas Martin",
    api_key: Optional[str] = None
) -> Dict[str, str]:
    """
    Génère deux formats de messages d'approche LinkedIn :
    1. Message de demande de connexion (strictement < 300 caractères).
    2. Message InMail direct complet (~100-150 mots).
    """
    company = job_details.get("Company", "l'Entreprise")
    title = job_details.get("Job Title", job_details.get("Title", "Business Analyst"))
    matched_str = ", ".join(matched_skills[:3]) if matched_skills else "Business Analysis"

    system_instruction = (
        "Tu es un spécialiste du réseautage professionnel sur LinkedIn en France. "
        "Tu écris des messages d'approche qui captent immédiatement l'attention des recruteurs et managers, "
        "avec un ton courtois, direct et sans lourdeur."
    )

    prompt = f"""
Rédige 2 messages d'approche LinkedIn en français pour contacter le recruteur ou manager du poste de **{title}** chez **{company}** :

- Nom du candidat : {candidate_name}
- Compétences clés alignées : {matched_str}

Retourne STRICTEMENT le format suivant avec les balises indiquées :

[CONNEXION_NOTE]
(Ici une note d'invitation percutante de MOINS DE 280 CARACTÈRES espaces compris pour la demande de connexion LinkedIn).
[/CONNEXION_NOTE]

[INMAIL_MESSAGE]
(Ici un message InMail complet de 100 à 150 mots : Salutations, accroche sur l'offre de {title}, 2 arguments percutants liant le profil à {company}, appel à l'action pour un bref échange de 10 minutes, formule de politesse).
[/INMAIL_MESSAGE]
"""

    response_text = call_gemini_api(prompt, system_instruction=system_instruction, api_key=api_key)

    # Extraction des sections
    conn_match = re.search(r'\[CONNEXION_NOTE\](.*?)\[/CONNEXION_NOTE\]', response_text, re.DOTALL)
    inmail_match = re.search(r'\[INMAIL_MESSAGE\](.*?)\[/INMAIL_MESSAGE\]', response_text, re.DOTALL)

    note = conn_match.group(1).strip() if conn_match else response_text[:280]
    inmail = inmail_match.group(1).strip() if inmail_match else response_text

    return {
        "connection_note": note,
        "inmail_message": inmail
    }


def suggest_ats_bullets(
    cv_text: str,
    job_details: Dict[str, Any],
    missing_skills: List[str],
    api_key: Optional[str] = None
) -> str:
    """
    Propose 3 gạch đầu dòng (bullet points) hành động theo phương pháp STAR/XYZ
    để tích hợp vào CV nhằm bù đắp các từ khóa ATS còn thiếu.
    """
    if not missing_skills:
        return "🎉 Votre profil couvre déjà l'intégralité des compétences clés de l'offre !"

    company = job_details.get("Company", "l'Entreprise")
    title = job_details.get("Job Title", job_details.get("Title", "Business Analyst"))
    missing_str = ", ".join(missing_skills)

    prompt = f"""
En tant qu'expert en recrutement et optimisation ATS pour le marché français, analyse les compétences manquantes suivantes :
**{missing_str}** pour le poste de **{title}** chez **{company}**.

Propose 3 puces de réalisations (bullet points) percutantes en français, rédigées au format STAR/XYZ (Verbe d'action + Contexte/Outil + Résultat chiffré), que le candidat peut adapter et intégrer dans ses expériences de CV pour franchir le filtre ATS :

Donne des exemples réalistes et concrets.
"""
    return call_gemini_api(prompt, api_key=api_key)
